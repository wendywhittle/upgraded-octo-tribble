from __future__ import annotations

import json
import math
import os
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook

STATUSES = ("NEW", "REVIEWING", "PURSUE", "HOLD", "PASS")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def database_path() -> str:
    return os.getenv("ALETHEIA_DB_PATH", str(Path("data") / "deals.db"))


def connect() -> sqlite3.Connection:
    path = Path(database_path())
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute(
        "CREATE TABLE IF NOT EXISTS deals (deal_id TEXT PRIMARY KEY, payload TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL)"
    )
    return conn


def finite(value: float | None) -> bool:
    return value is not None and math.isfinite(value)


def mortgage_payment(principal: float, annual_rate: float, years: float) -> float:
    months = years * 12
    monthly_rate = annual_rate / 12
    if monthly_rate == 0:
        return principal / months
    factor = (1 + monthly_rate) ** months
    return principal * monthly_rate * factor / (factor - 1)


def npv(rate: float, cashflows: list[float]) -> float:
    return sum(cf / ((1 + rate) ** i) for i, cf in enumerate(cashflows))


def irr(cashflows: list[float]) -> float | None:
    if not cashflows or not (any(x < 0 for x in cashflows) and any(x > 0 for x in cashflows)):
        return None
    low, high = -0.9999, 10.0
    if npv(low, cashflows) * npv(high, cashflows) > 0:
        return None
    for _ in range(200):
        mid = (low + high) / 2
        value = npv(mid, cashflows)
        if abs(value) < 1e-10:
            return mid
        if npv(low, cashflows) * value <= 0:
            high = mid
        else:
            low = mid
    return (low + high) / 2


def underwrite(deal: dict[str, Any]) -> dict[str, Any]:
    price = deal.get("purchase_price")
    explicit_noi = deal.get("noi")
    egi = deal.get("egi")
    expenses = deal.get("operating_expenses")
    ltv = deal.get("ltv")
    rate = deal.get("interest_rate")
    amort = deal.get("amortization_years")
    hold = deal.get("hold_years")
    exit_cap = deal.get("exit_cap_rate")
    closing_costs = deal.get("closing_costs") or 0

    derived: dict[str, Any] = {}
    missing: list[str] = []

    if finite(explicit_noi):
        noi = explicit_noi
        noi_basis = "supplied"
    elif finite(egi) and finite(expenses):
        noi = egi - expenses
        noi_basis = "calculated_from_egi_minus_operating_expenses"
        derived["noi"] = noi
        derived["noi_basis"] = noi_basis
    else:
        noi = None
        noi_basis = "unavailable"

    if not finite(price) or price <= 0:
        missing.append("purchase_price")
    if not finite(noi):
        missing.append("noi")

    if finite(price) and price > 0 and finite(noi):
        derived["cap_rate"] = noi / price
        derived["cap_rate_basis"] = noi_basis

    loan = None
    annual_debt = None
    initial_equity = None

    if finite(price) and price > 0 and finite(ltv):
        loan = price * ltv
        initial_equity = price - loan + closing_costs
        derived["loan_amount"] = loan
        derived["initial_equity"] = initial_equity
    elif ltv is None:
        missing.append("ltv")

    if loan is not None and finite(rate) and finite(amort):
        annual_debt = mortgage_payment(loan, rate, amort) * 12
        derived["annual_debt_service"] = annual_debt
    elif loan is not None:
        if rate is None:
            missing.append("interest_rate")
        if amort is None:
            missing.append("amortization_years")

    if finite(noi) and finite(annual_debt) and annual_debt > 0:
        derived["dscr"] = noi / annual_debt
        derived["annual_cash_flow"] = noi - annual_debt
        if finite(initial_equity) and initial_equity > 0:
            derived["cash_on_cash"] = (noi - annual_debt) / initial_equity

    if finite(noi) and finite(exit_cap) and exit_cap > 0:
        derived["exit_value"] = noi / exit_cap
    elif exit_cap is None:
        missing.append("exit_cap_rate")

    if finite(hold) and hold > 0 and "exit_value" in derived and finite(initial_equity) and initial_equity > 0:
        annual_cash = derived.get("annual_cash_flow")
        if finite(annual_cash):
            final_cash = annual_cash + derived["exit_value"] - (loan or 0)
            cashflows = [-initial_equity] + [annual_cash] * (int(hold) - 1) + [final_cash]
            result = irr(cashflows)
            if result is not None:
                derived["irr"] = result
                derived["equity_multiple"] = sum(max(x, 0) for x in cashflows) / initial_equity
    elif hold is None:
        missing.append("hold_years")

    return {"derived": derived, "missing": sorted(set(missing))}


def next_deal_id(conn: sqlite3.Connection) -> str:
    year = datetime.now(timezone.utc).year
    row = conn.execute("SELECT COUNT(*) AS n FROM deals").fetchone()
    number = int(row["n"]) + 1
    while conn.execute("SELECT 1 FROM deals WHERE deal_id = ?", (f"AT-{year}-{number:06d}",)).fetchone():
        number += 1
    return f"AT-{year}-{number:06d}"


def create_deal(payload: dict[str, Any]) -> dict[str, Any]:
    conn = connect()
    try:
        now = utc_now()
        deal_id = next_deal_id(conn)
        result = underwrite(payload)
        record = {
            "deal_id": deal_id,
            "original_inputs": payload,
            "derived": result["derived"],
            "missing": result["missing"],
            "status": "NEW",
            "created_at": now,
            "updated_at": now,
        }
        conn.execute(
            "INSERT INTO deals VALUES (?, ?, ?, ?)",
            (deal_id, json.dumps(record), now, now),
        )
        conn.commit()
        return record
    finally:
        conn.close()


def list_deals() -> list[dict[str, Any]]:
    conn = connect()
    try:
        return [json.loads(row["payload"]) for row in conn.execute("SELECT payload FROM deals ORDER BY created_at DESC")]
    finally:
        conn.close()


def get_deal(deal_id: str) -> dict[str, Any] | None:
    conn = connect()
    try:
        row = conn.execute("SELECT payload FROM deals WHERE deal_id = ?", (deal_id,)).fetchone()
        return json.loads(row["payload"]) if row else None
    finally:
        conn.close()


def update_status(deal_id: str, status: str) -> dict[str, Any] | None:
    if status not in STATUSES:
        raise ValueError("Invalid pipeline status")
    conn = connect()
    try:
        row = conn.execute("SELECT payload FROM deals WHERE deal_id = ?", (deal_id,)).fetchone()
        if not row:
            return None
        record = json.loads(row["payload"])
        record["status"] = status
        record["updated_at"] = utc_now()
        conn.execute(
            "UPDATE deals SET payload = ?, updated_at = ? WHERE deal_id = ?",
            (json.dumps(record), record["updated_at"], deal_id),
        )
        conn.commit()
        return record
    finally:
        conn.close()


def build_excel(record: dict[str, Any]) -> bytes:
    from io import BytesIO

    wb = Workbook()
    summary = wb.active
    summary.title = "Deal Summary"
    summary.append(["AletheiaTelos Deal Flow Engine"])
    summary.append(["Deal ID", record["deal_id"]])
    summary.append(["Status", record["status"]])
    summary.append([])
    for key, value in record["original_inputs"].items():
        summary.append([key.replace("_", " ").title(), value])
    summary.append([])
    summary.append(["Calculated", "Value"])
    for key, value in record["derived"].items():
        summary.append([key.replace("_", " ").title(), value])

    for title, keys in [
        ("Operating Inputs", ["egi", "operating_expenses", "occupancy", "noi"]),
        ("Financing", ["ltv", "interest_rate", "amortization_years", "closing_costs"]),
        ("Returns", ["hold_years", "exit_cap_rate"]),
    ]:
        sheet = wb.create_sheet(title)
        sheet.append(["Field", "Value"])
        for key in keys:
            sheet.append([key, record["original_inputs"].get(key)])

    missing = wb.create_sheet("Missing Data")
    missing.append(["Missing / Unavailable Information"])
    for item in record["missing"]:
        missing.append([item])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
